"""将三份 Excel 知识库转换为可追溯的前端 JSON 数据。"""
from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
LOCAL_PACKAGES = ROOT / ".python-packages"
if LOCAL_PACKAGES.exists():
    sys.path.insert(0, str(LOCAL_PACKAGES))

from openpyxl import load_workbook  # noqa: E402

OUTPUT = ROOT / "src" / "data"
FILES = ["十篇核心论文.xlsx", "五类环境指标.xlsx", "秦岭植物园和牛背梁分段数据.xlsx"]


def clean(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def read_sheet(file_name: str, sheet_name: str) -> list[dict[str, Any]]:
    ws = load_workbook(ROOT / file_name, data_only=True, read_only=True)[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(clean(value) or f"未命名字段_{index + 1}") for index, value in enumerate(rows[0])]
    records: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows[1:], start=2):
        values = [clean(value) for value in row]
        if not any(value not in (None, "") for value in values):
            continue
        record = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        record["_source"] = {"file": file_name, "sheet": sheet_name, "row": row_index}
        records.append(record)
    return records


def raw_workbooks() -> tuple[list[dict[str, Any]], int]:
    workbooks: list[dict[str, Any]] = []
    total = 0
    for file_name in FILES:
        workbook = load_workbook(ROOT / file_name, data_only=True, read_only=True)
        sheets = []
        for ws in workbook.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(clean(value) or f"未命名字段_{index + 1}") for index, value in enumerate(rows[0])]
            data = []
            for row_index, row in enumerate(rows[1:], start=2):
                values = [clean(value) for value in row]
                if not any(value not in (None, "") for value in values):
                    continue
                data.append({"row": row_index, "values": {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}})
            total += len(data)
            sheets.append({"name": ws.title, "headers": headers, "recordCount": len(data), "records": data})
        workbooks.append({"file": file_name, "sheetCount": len(sheets), "sheets": sheets})
    return workbooks, total


def write_json(name: str, data: Any) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    missing = [name for name in FILES if not (ROOT / name).exists()]
    if missing:
        raise FileNotFoundError(f"缺少 Excel 文件: {', '.join(missing)}")

    raw, raw_count = raw_workbooks()
    literature = read_sheet("十篇核心论文.xlsx", "10篇核心文献")

    population_rules = read_sheet("十篇核心论文.xlsx", "六类人群规则")
    population_match = {row["人群"]: row for row in read_sheet("五类环境指标.xlsx", "六类人群匹配")}
    populations = []
    for row in population_rules:
        companion = population_match.get(row["用户类型"], {})
        populations.append({**row, "补充匹配": {key: value for key, value in companion.items() if key != "_source"}, "_matchSource": companion.get("_source")})

    indicators = [
        *[{**row, "记录类型": "五类指标标准"} for row in read_sheet("五类环境指标.xlsx", "五类指标标准")],
        *[{**row, "记录类型": "变量字典"} for row in read_sheet("十篇核心论文.xlsx", "变量字典")],
        *[{**row, "记录类型": "分段参数规则"} for row in read_sheet("秦岭植物园和牛背梁分段数据.xlsx", "五类参数规则")],
    ]

    rules = [
        *[{**row, "规则集": "筛选逻辑"} for row in read_sheet("五类环境指标.xlsx", "Coze筛选逻辑")],
        *[{**row, "规则集": "知识库约束"} for row in read_sheet("十篇核心论文.xlsx", "Coze导入说明")],
        *[{**row, "规则集": "分段执行顺序"} for row in read_sheet("秦岭植物园和牛背梁分段数据.xlsx", "Coze执行顺序")],
    ]

    sites = [
        *[{**row, "记录类型": "具体分段"} for row in read_sheet("秦岭植物园和牛背梁分段数据.xlsx", "区域分段数据库")],
        *[{**row, "记录类型": "场地候选"} for row in read_sheet("十篇核心论文.xlsx", "秦岭场地候选")],
        *[{**row, "记录类型": "具体区域"} for row in read_sheet("五类环境指标.xlsx", "秦岭具体区域")],
    ]

    sources = read_sheet("五类环境指标.xlsx", "文献与官方来源")
    field_tasks = [
        *[{**row, "清单类型": "现场采集"} for row in read_sheet("五类环境指标.xlsx", "现场采集清单")],
        *[{**row, "清单类型": "实地调查"} for row in read_sheet("秦岭植物园和牛背梁分段数据.xlsx", "实地调查任务")],
    ]

    datasets = {
        "literature.json": literature,
        "populations.json": populations,
        "environment-indicators.json": indicators,
        "screening-rules.json": rules,
        "qinling-sites.json": sites,
        "sources.json": sources,
        "field-tasks.json": field_tasks,
        "raw-workbooks.json": raw,
    }
    for name, data in datasets.items():
        write_json(name, data)

    manifest = {
        "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "sourceFiles": FILES,
        "workbookCount": len(FILES),
        "sheetCount": sum(book["sheetCount"] for book in raw),
        "rawRecordCount": raw_count,
        "datasets": {name: len(data) for name, data in datasets.items() if isinstance(data, list)},
        "integrity": {
            "missingSourceFiles": missing,
            "unknownValuesPreserved": True,
            "numericThresholdsInferred": False,
        },
    }
    write_json("import-manifest.json", manifest)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

